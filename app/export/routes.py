"""XLSX export route: /api/export/xlsx."""
from __future__ import annotations

import io
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth.dependencies import get_current_user
from app.db.models import FoodEntry, User
from app.db.session import get_db

router = APIRouter(prefix="/api/export", tags=["export"])


@router.get("/xlsx")
def export_xlsx(
    start_date: Optional[date] = None, end_date: Optional[date] = None,
    user: User = Depends(get_current_user), db: Session = Depends(get_db),
):
    """Kullanıcının verilerini XLSX olarak indir. Google Sheets'e import edilebilir."""
    stmt = select(FoodEntry).where(FoodEntry.user_id == user.id)
    if start_date:
        stmt = stmt.where(FoodEntry.entry_date >= start_date)
    if end_date:
        stmt = stmt.where(FoodEntry.entry_date <= end_date)
    stmt = stmt.order_by(FoodEntry.entry_date.asc(), FoodEntry.entry_time.asc())
    rows = list(db.scalars(stmt))

    wb = Workbook()
    ws = wb.active
    ws.title = "food_log"
    headers = [
        "id", "date", "time", "meal_type", "food_name", "weight_g",
        "calories", "protein_g", "carbs_g", "fat_g",
        "fiber_g", "sugar_g", "sodium_mg", "notes", "created_at",
    ]
    ws.append(headers)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1A1410")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
    for e in rows:
        ws.append([
            str(e.id), e.entry_date.isoformat(), e.entry_time or "",
            e.meal_type, e.food_name, e.weight_g,
            e.calories, e.protein_g, e.carbs_g, e.fat_g,
            e.fiber_g, e.sugar_g, e.sodium_mg,
            e.notes or "", e.created_at.replace(tzinfo=None).isoformat(timespec="seconds"),
        ])
    for i, h in enumerate(headers, 1):
        col = ws.column_dimensions[chr(64 + i) if i <= 26 else "AA"]
        col.width = max(12, len(h) + 2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"food-log-{user.email.split('@')[0]}-{date.today().isoformat()}.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
